import os
from docx import Document
from openpyxl import Workbook
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

os.makedirs("generated", exist_ok=True)

def create_word_doc(filename: str, content: str) -> str:
    path = os.path.join("generated", filename)
    doc = Document()
    doc.add_heading("NPICADO_IA - Documento Generado", level=1)
    for line in content.split("\n"):
        if line.strip():
            doc.add_paragraph(line)
    doc.save(path)
    return path

def create_excel_doc(filename: str, content: str) -> str:
    path = os.path.join("generated", filename)
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte NPICADO"
    
    lines = content.split("\n")
    for row_idx, line in enumerate(lines, start=1):
        ws.cell(row=row_idx, column=1, value=line)
        
    wb.save(path)
    return path

def create_pdf_doc(filename: str, content: str) -> str:
    path = os.path.join("generated", filename)
    c = canvas.Canvas(path, pagesize=letter)
    y = 750
    c.setFont("Helvetica-Bold", 14)
    c.drawString(50, y, "NPICADO_IA - Documento Generado")
    y -= 30
    c.setFont("Helvetica", 10)
    
    for line in content.split("\n"):
        if y < 50:
            c.showPage()
            y = 750
            c.setFont("Helvetica", 10)
        c.drawString(50, y, line[:90])
        y -= 15
        
    c.save()
    return path